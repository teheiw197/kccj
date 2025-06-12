"""
课程表解析模块
支持多种格式：Word、Excel、图片、纯文本
"""
from typing import List, Dict, Optional
import docx
import aiohttp
import re
import os
import json
import openpyxl
import logging
from PIL import Image
import io

logger = logging.getLogger(__name__)

def parse_word(file_path: str) -> List[Dict]:
    """解析Word课程表"""
    try:
        result = []
        doc = docx.Document(file_path)
        for table in doc.tables:
            for row in table.rows[1:]:  # 跳过表头
                cells = [cell.text.strip() for cell in row.cells]
                if len(cells) >= 4:
                    result.append({
                        "weekday": extract_weekday(cells[1]),
                        "time": cells[1],
                        "course": cells[0],
                        "classroom": cells[2],
                        "teacher": cells[3]
                    })
        return result
    except Exception as e:
        logger.error(f"解析Word文件失败: {str(e)}")
        return []

def parse_xlsx(file_path: str) -> List[Dict]:
    """解析Excel课程表"""
    try:
        result = []
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue  # 跳过表头
            cells = [str(cell).strip() if cell is not None else '' for cell in row]
            if len(cells) >= 4:
                result.append({
                    "weekday": extract_weekday(cells[1]),
                    "time": cells[1],
                    "course": cells[0],
                    "classroom": cells[2],
                    "teacher": cells[3]
                })
        return result
    except Exception as e:
        logger.error(f"解析Excel文件失败: {str(e)}")
        return []

async def parse_image(file_path: str, ocr_api_url: str, ocr_api_key: str = None) -> List[Dict]:
    """通过OCR识别图片课程表"""
    try:
        # 读取图片为二进制
        with open(file_path, "rb") as f:
            img_bytes = f.read()
            
        headers = {}
        if ocr_api_key:
            headers["Authorization"] = ocr_api_key
            
        data = aiohttp.FormData()
        data.add_field('image', img_bytes, 
                      filename=os.path.basename(file_path), 
                      content_type='application/octet-stream')
                      
        async with aiohttp.ClientSession() as session:
            async with session.post(ocr_api_url, headers=headers, data=data) as resp:
                resp_json = await resp.json()
                text = resp_json.get("text") or resp_json.get("data", {}).get("text", "")
                
        return parse_text_schedule(text)
    except Exception as e:
        logger.error(f"解析图片失败: {str(e)}")
        return []

def parse_text_schedule(text_content: str) -> List[Dict]:
    """解析纯文本课程表"""
    try:
        result = []
        lines = [line.strip() for line in text_content.split('\n') if line.strip()]
        
        for line in lines:
            # 尝试多种格式匹配
            patterns = [
                # 格式1：课程名 星期 时间 教室 教师
                r'(.+?)\s+(周[一二三四五六日])\s+(第\d+-\d+节)\s+(.+?)\s+(.+)',
                # 格式2：星期 时间 课程名 教室 教师
                r'(周[一二三四五六日])\s+(第\d+-\d+节)\s+(.+?)\s+(.+?)\s+(.+)',
                # 格式3：课程名 星期时间 教室 教师
                r'(.+?)\s+(周[一二三四五六日]第\d+-\d+节)\s+(.+?)\s+(.+)',
            ]
            
            for pattern in patterns:
                m = re.match(pattern, line)
                if m:
                    groups = m.groups()
                    if len(groups) == 5:  # 格式1或2
                        if "课程" in groups[0]:  # 格式1
                            result.append({
                                "weekday": groups[1],
                                "time": groups[2],
                                "course": groups[0],
                                "classroom": groups[3],
                                "teacher": groups[4]
                            })
                        else:  # 格式2
                            result.append({
                                "weekday": groups[0],
                                "time": groups[1],
                                "course": groups[2],
                                "classroom": groups[3],
                                "teacher": groups[4]
                            })
                    elif len(groups) == 4:  # 格式3
                        result.append({
                            "weekday": extract_weekday(groups[1]),
                            "time": groups[1],
                            "course": groups[0],
                            "classroom": groups[2],
                            "teacher": groups[3]
                        })
                    break
                    
        return result
    except Exception as e:
        logger.error(f"解析文本失败: {str(e)}")
        return []

def extract_weekday(time_str: str) -> str:
    """从时间字符串中提取星期信息"""
    weekday_pattern = r'周[一二三四五六日]'
    match = re.search(weekday_pattern, time_str)
    return match.group(0) if match else ""

def get_class_time(time_str: str) -> Optional[tuple]:
    """从时间字符串中提取具体时间"""
    try:
        # 处理"第1-2节"格式
        if "节" in time_str:
            section = int(time_str.split("-")[0].replace("第", ""))
            section_times = {
                1: (8, 0), 2: (8, 55), 3: (10, 0), 4: (10, 55),
                5: (14, 0), 6: (14, 55), 7: (16, 0), 8: (16, 55),
                9: (19, 0), 10: (19, 55), 11: (20, 50)
            }
            return section_times.get(section)
            
        # 处理"8:00"格式
        if ":" in time_str:
            hour, minute = map(int, time_str.split(":"))
            return (hour, minute)
            
        return None
    except Exception as e:
        logger.error(f"解析时间失败: {str(e)}")
        return None 